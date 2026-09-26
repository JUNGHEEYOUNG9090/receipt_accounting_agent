import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


OUTPUT_FILE = "receipt_comparison.xlsx"


def create_excel(result):

    wb = Workbook()

    # ========================================
    # 결과 데이터 매핑
    # ========================================

    ocr_map = {
        item["filename"]: item
        for item in result["ocr_results"]
    }

    llm_map = {
        item["filename"]: item
        for item in result["receipts"]
    }

    vision_map = {
        item["filename"]: item
        for item in result.get("vision_results", [])
    }

    # ========================================
    # 1. 비교결과
    # ========================================

    ws = wb.active
    ws.title = "비교결과"

    ws.append([
        "파일명",
        "필드",
        "실제값",
        "OCR 결과",
        "LLM 결과",
        "Vision 결과",
    ])

    for filename in ocr_map:

        ocr_data = ocr_map[filename]

        ocr_text = ocr_data.get(
            "ocr_text",
            ""
        )

        llm_data = llm_map.get(
            filename,
            {}
        )

        llm_receipt = llm_data.get(
            "receipt",
            {}
        )

        vision_data = vision_map.get(
            filename,
            {}
        )

        vision_receipt = vision_data.get(
            "receipt",
            {}
        )

        # LLM JSON 전체
        llm_text = ""

        if llm_receipt:
            llm_text = json.dumps(
                llm_receipt,
                ensure_ascii=False,
                indent=2
            )

        # Vision JSON 전체
        vision_text = ""

        if vision_receipt:
            vision_text = json.dumps(
                vision_receipt,
                ensure_ascii=False,
                indent=2
            )

        # 현재는 Ground Truth가 없으므로 빈칸
        actual_value = ""

        ws.append([
            filename,
            "OCR 전체 결과",
            actual_value,
            ocr_text,
            llm_text,
            vision_text,
        ])

    # ========================================
    # 2. 성능
    # ========================================

    perf = wb.create_sheet("성능")

    perf.append([
        "파일명",
        "OCR 시간(초)",
        "LLM 시간(초)",
        "Vision 시간(초)",
        "전체 시간(초)",
        "LLM 비용",
        "Vision 비용",
        "Vision 호출 여부",
    ])

    for filename in ocr_map:

        ocr_data = ocr_map[filename]

        ocr_time = ocr_data.get(
            "ocr_time",
            ""
        )

        llm_data = llm_map.get(
            filename,
            {}
        )

        vision_data = vision_map.get(
            filename,
            {}
        )

        llm_time = llm_data.get(
            "llm_time",
            ""
        )

        vision_time = vision_data.get(
            "vision_time",
            ""
        )

        # OCR + LLM + Vision
        times = [
            value
            for value in [
                ocr_time,
                llm_time,
                vision_time,
            ]
            if value != ""
        ]

        if times:
            total_time = round(
                sum(times),
                3
            )
        else:
            total_time = ""

        vision_called = (
            "O"
            if vision_data
            else "X"
        )

        perf.append([
            filename,
            ocr_time,
            llm_time,
            vision_time,
            total_time,
            llm_data.get(
                "llm_cost",
                ""
            ),
            vision_data.get(
                "vision_cost",
                ""
            ),
            vision_called,
        ])

    # ========================================
    # 3. 테스트 조건
    # ========================================

    condition_ws = wb.create_sheet(
        "테스트조건"
    )

    condition_ws.append([
        "파일명",
        "촬영 조건",
    ])

    # 현재는 별도 조건 데이터가 없으므로
    # 이미지 파일만 등록
    for filename in ocr_map:

        condition_ws.append([
            filename,
            "",
        ])

    # ========================================
    # 스타일
    # ========================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for sheet in wb.worksheets:

        sheet.freeze_panes = "A2"

        for cell in sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row in sheet.iter_rows():

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    # ========================================
    # 비교결과 열 너비
    # ========================================

    comparison_widths = {
        "A": 24,
        "B": 18,
        "C": 70,
        "D": 70,
        "E": 70,
        "F": 70,
    }

    for column, width in comparison_widths.items():

        ws.column_dimensions[column].width = width

    # ========================================
    # 성능 열 너비
    # ========================================

    perf_widths = {
        "A": 24,
        "B": 16,
        "C": 16,
        "D": 18,
        "E": 16,
        "F": 16,
        "G": 16,
        "H": 18,
    }

    for column, width in perf_widths.items():

        perf.column_dimensions[column].width = width

    # ========================================
    # 테스트 조건 열 너비
    # ========================================

    condition_ws.column_dimensions["A"].width = 24
    condition_ws.column_dimensions["B"].width = 60

    # ========================================
    # 필터
    # ========================================

    ws.auto_filter.ref = ws.dimensions
    perf.auto_filter.ref = perf.dimensions
    condition_ws.auto_filter.ref = condition_ws.dimensions

    # ========================================
    # 저장
    # ========================================

    wb.save(OUTPUT_FILE)

    print(
        f"\nExcel 생성 완료 : {OUTPUT_FILE}"
    )
